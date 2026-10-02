# -*- coding: utf-8 -*-
"""
PassManager Professional Edition - Flask Backend
Argon2id KDF + AES-256-GCM encryption + Encrypted JSON vault (no database)
"""
import os, sys, json, base64, secrets, time, csv, io, re, uuid, string
import hashlib, threading
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path

from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

try:
    from argon2.low_level import hash_secret_raw, Type
    HAS_ARGON2 = True
except ImportError:
    HAS_ARGON2 = False

try:
    from openpyxl import Workbook, load_workbook
    HAS_EXCEL = True
except ImportError:
    HAS_EXCEL = False

try:
    import pyodbc
    HAS_MDB = True
except ImportError:
    HAS_MDB = False

# ================= CONFIG =================
BASE_DIR = Path(__file__).parent.resolve()
VAULT_PATH = BASE_DIR / "vault.json"
BACKUP_DIR = BASE_DIR / "backups"
AUDIT_PATH = BASE_DIR / "audit.log"
BACKUP_DIR.mkdir(exist_ok=True)

LOCK_TIMEOUT = 300
MAX_ATTEMPTS = 5
LOCKOUT_SECONDS = 300
CLIPBOARD_CLEAR_SECONDS = 20
BACKUP_KEEP = 10

# ================= APP =================
app = Flask(__name__)
app.secret_key = secrets.token_hex(32)
CORS(app, supports_credentials=True)

STATE = {
    "unlocked": False,
    "key": None,
    "salt": None,
    "data": None,
    "created": None,
    "last_activity": 0,
    "failed": {},
}
STATE_LOCK = threading.Lock()

# ================= HELPERS =================
def now_iso(): return datetime.utcnow().isoformat()

def log_audit(event, detail=""):
    try:
        ip = request.remote_addr if request else "-"
        with open(AUDIT_PATH, "a", encoding="utf-8") as f:
            f.write(f"{now_iso()} | {ip} | {event} | {detail}\n")
    except Exception:
        pass

def derive_key(password, salt):
    if HAS_ARGON2:
        return hash_secret_raw(
            secret=password.encode("utf-8"), salt=salt,
            time_cost=3, memory_cost=65536, parallelism=4,
            hash_len=32, type=Type.ID,
        )
    return hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 200000, 32)

def enc(key, data):
    nonce = secrets.token_bytes(12)
    return nonce, AESGCM(key).encrypt(nonce, data, None)

def dec(key, nonce, ct):
    return AESGCM(key).decrypt(nonce, ct, None)

def atomic_write(path: Path, content: str):
    tmp = path.with_name(path.name + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(content); f.flush(); os.fsync(f.fileno())
    os.replace(tmp, path)

def save_vault():
    with STATE_LOCK:
        if not STATE["unlocked"]: return
        data_json = json.dumps(STATE["data"], ensure_ascii=False).encode("utf-8")
        nonce, ct = enc(STATE["key"], data_json)
        checksum = hashlib.sha256(ct).hexdigest()
        vault = {
            "version": 1,
            "kdf": "argon2id" if HAS_ARGON2 else "pbkdf2-hmac-sha256",
            "salt": base64.b64encode(STATE["salt"]).decode(),
            "nonce": base64.b64encode(nonce).decode(),
            "ciphertext": base64.b64encode(ct).decode(),
            "checksum": checksum,
            "created": STATE["created"],
            "modified": now_iso(),
        }
        atomic_write(VAULT_PATH, json.dumps(vault, indent=2))
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        atomic_write(BACKUP_DIR / f"vault_{ts}.json", json.dumps(vault, indent=2))
        backups = sorted(BACKUP_DIR.glob("vault_*.json"))
        for b in backups[:-BACKUP_KEEP]:
            try: b.unlink()
            except: pass
        STATE["last_activity"] = time.time()

def check_lock_timeout():
    if STATE["unlocked"] and time.time() - STATE["last_activity"] > LOCK_TIMEOUT:
        lock_vault("timeout"); return True
    return False

def lock_vault(reason="manual"):
    with STATE_LOCK:
        STATE["unlocked"] = False
        STATE["key"] = None
        STATE["data"] = None
    log_audit("LOCK", reason)

def require_unlock(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if check_lock_timeout():
            return jsonify({"error": "Vault locked (inactivity)"}), 401
        if not STATE["unlocked"]:
            return jsonify({"error": "Vault is locked"}), 401
        STATE["last_activity"] = time.time()
        return f(*args, **kwargs)
    return wrapper

def password_strength(pw):
    if not pw: return 0
    s = min(40, len(pw) * 3)
    if re.search(r"[a-z]", pw): s += 10
    if re.search(r"[A-Z]", pw): s += 15
    if re.search(r"\d", pw): s += 15
    if re.search(r"[^A-Za-z0-9]", pw): s += 20
    return min(100, s)

# ================= ROUTES =================
@app.route("/api/health", methods=["GET"])
def health():
    return jsonify({"status":"ok","argon2":HAS_ARGON2,"excel":HAS_EXCEL,"mdb":HAS_MDB})

@app.route("/api/status", methods=["GET"])
def status():
    check_lock_timeout()
    return jsonify({
        "unlocked": STATE["unlocked"],
        "exists": VAULT_PATH.exists(),
        "lock_timeout": LOCK_TIMEOUT,
    })

@app.route("/api/unlock", methods=["POST"])
def unlock():
    ip = request.remote_addr or "?"
    fdata = STATE["failed"].get(ip, [0, 0])
    if fdata[1] > time.time():
        wait = int(fdata[1] - time.time())
        return jsonify({"error": f"Too many attempts. Wait {wait}s"}), 429

    body = request.get_json() or {}
    password = body.get("password", "")
    if not password:
        return jsonify({"error": "Password required"}), 400

    if not VAULT_PATH.exists():
        salt = secrets.token_bytes(16)
        key = derive_key(password, salt)
        with STATE_LOCK:
            STATE["unlocked"] = True; STATE["key"] = key; STATE["salt"] = salt
            STATE["data"] = []; STATE["created"] = now_iso()
            STATE["last_activity"] = time.time()
        save_vault()
        log_audit("VAULT_CREATED", ip)
        return jsonify({"success": True, "created": True})

    try:
        vault = json.loads(VAULT_PATH.read_text(encoding="utf-8"))
    except Exception:
        return jsonify({"error": "Vault corrupted"}), 500

    salt = base64.b64decode(vault["salt"])
    nonce = base64.b64decode(vault["nonce"])
    ct = base64.b64decode(vault["ciphertext"])
    if hashlib.sha256(ct).hexdigest() != vault.get("checksum"):
        return jsonify({"error": "Integrity check failed"}), 500

    key = derive_key(password, salt)
    try:
        plain = dec(key, nonce, ct)
    except Exception:
        fdata = STATE["failed"].get(ip, [0, 0])
        fdata[0] += 1
        if fdata[0] >= MAX_ATTEMPTS:
            fdata[1] = time.time() + LOCKOUT_SECONDS; fdata[0] = 0
        STATE["failed"][ip] = fdata
        log_audit("FAILED_UNLOCK", ip)
        return jsonify({"error": "Invalid master password"}), 401

    STATE["failed"][ip] = [0, 0]
    data = json.loads(plain.decode("utf-8"))
    with STATE_LOCK:
        STATE["unlocked"] = True; STATE["key"] = key; STATE["salt"] = salt
        STATE["data"] = data; STATE["created"] = vault.get("created", now_iso())
        STATE["last_activity"] = time.time()
    log_audit("UNLOCK", ip)
    return jsonify({"success": True})

@app.route("/api/lock", methods=["POST"])
def lock():
    lock_vault("api"); return jsonify({"success": True})

@app.route("/api/entries", methods=["GET"])
@require_unlock
def get_entries():
    return jsonify({"entries": STATE["data"]})

@app.route("/api/entries", methods=["POST"])
@require_unlock
def add_entry():
    b = request.get_json() or {}
    e = {
        "id": str(uuid.uuid4()),
        "site": (b.get("site") or "").strip(),
        "username": (b.get("username") or "").strip(),
        "password": b.get("password") or "",
        "url": (b.get("url") or "").strip(),
        "notes": b.get("notes") or "",
        "category": (b.get("category") or "general").strip() or "general",
        "created": now_iso(), "modified": now_iso(), "password_changed": now_iso(),
    }
    # ============ UPDATED: Kam se kam ek field bhara hona chahiye ============
    if not any([e["site"], e["username"], e["password"], e["url"], e["notes"]]):
        return jsonify({"error": "At least one field is required"}), 400
    # =========================================================================
    STATE["data"].append(e); save_vault()
    log_audit("ADD_ENTRY", e["site"] or e["username"] or e["url"] or "(untitled)")
    return jsonify({"success": True, "entry": e})

@app.route("/api/entries/<eid>", methods=["PUT"])
@require_unlock
def edit_entry(eid):
    b = request.get_json() or {}
    for e in STATE["data"]:
        if e["id"] == eid:
            old_pw = e.get("password")
            for k in ("site","username","password","url","notes","category"):
                if k in b: e[k] = b[k]
            e["modified"] = now_iso()
            if old_pw != e.get("password"):
                e["password_changed"] = now_iso()
            save_vault(); log_audit("EDIT_ENTRY", e["site"] or e["username"] or "(untitled)")
            return jsonify({"success": True, "entry": e})
    return jsonify({"error": "Not found"}), 404

@app.route("/api/entries/<eid>", methods=["DELETE"])
@require_unlock
def delete_entry(eid):
    for i, e in enumerate(STATE["data"]):
        if e["id"] == eid:
            STATE["data"].pop(i); save_vault()
            log_audit("DELETE_ENTRY", e["site"] or e["username"] or "(untitled)")
            return jsonify({"success": True})
    return jsonify({"error": "Not found"}), 404

@app.route("/api/generate-password", methods=["GET"])
def gen_pw():
    length = max(4, min(128, int(request.args.get("length", 16))))
    up = request.args.get("upper", "1") == "1"
    lo = request.args.get("lower", "1") == "1"
    dg = request.args.get("digits", "1") == "1"
    sy = request.args.get("symbols", "1") == "1"
    chars = ""
    if up: chars += string.ascii_uppercase
    if lo: chars += string.ascii_lowercase
    if dg: chars += string.digits
    if sy: chars += "!@#$%^&*()-_=+[]{};:,.<>?/|"
    if not chars: chars = string.ascii_letters + string.digits
    return jsonify({"password": "".join(secrets.choice(chars) for _ in range(length))})

@app.route("/api/change-master", methods=["POST"])
@require_unlock
def change_master():
    b = request.get_json() or {}
    new_pw = b.get("new_password", "")
    if len(new_pw) < 4: return jsonify({"error": "Too short"}), 400
    salt = secrets.token_bytes(16)
    key = derive_key(new_pw, salt)
    with STATE_LOCK:
        STATE["salt"] = salt; STATE["key"] = key
    save_vault(); log_audit("CHANGE_MASTER")
    return jsonify({"success": True})

# ================= EXPORT / IMPORT =================
def entries_to_rows(entries):
    return [{
        "site": e.get("site",""), "username": e.get("username",""),
        "password": e.get("password",""), "url": e.get("url",""),
        "notes": e.get("notes",""), "category": e.get("category",""),
    } for e in entries]

@app.route("/api/export", methods=["POST"])
@require_unlock
def export_data():
    b = request.get_json() or {}
    fmt = (b.get("format") or "json").lower()
    rows = entries_to_rows(STATE["data"])
    log_audit("EXPORT", fmt)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")

    if fmt == "json":
        buf = io.BytesIO(json.dumps(rows, indent=2, ensure_ascii=False).encode("utf-8"))
        return send_file(buf, mimetype="application/json",
                         as_attachment=True, download_name=f"passmanager_{ts}.json")
    if fmt == "csv":
        sio = io.StringIO()
        w = csv.DictWriter(sio, fieldnames=["site","username","password","url","notes","category"])
        w.writeheader()
        for r in rows: w.writerow(r)
        buf = io.BytesIO(sio.getvalue().encode("utf-8-sig"))
        return send_file(buf, mimetype="text/csv",
                         as_attachment=True, download_name=f"passmanager_{ts}.csv")
    if fmt in ("xls","xlsx"):
        if not HAS_EXCEL: return jsonify({"error": "openpyxl missing"}), 500
        wb = Workbook(); ws = wb.active; ws.title = "Passwords"
        ws.append(["site","username","password","url","notes","category"])
        for r in rows:
            ws.append([r["site"],r["username"],r["password"],r["url"],r["notes"],r["category"]])
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        return send_file(buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True, download_name=f"passmanager_{ts}.xlsx")
    if fmt == "mdb":
        return jsonify({"error": "MDB export needs MS Access driver. Use XLSX/CSV."}), 400
    return jsonify({"error": "Unsupported"}), 400

@app.route("/api/import", methods=["POST"])
@require_unlock
def import_data():
    if "file" not in request.files: return jsonify({"error": "No file"}), 400
    f = request.files["file"]
    name = f.filename or ""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    rows = []
    try:
        content = f.read()
        if ext == "json":
            parsed = json.loads(content.decode("utf-8"))
            if isinstance(parsed, dict) and "entries" in parsed:
                parsed = parsed["entries"]
            rows = parsed
        elif ext == "csv":
            try: text = content.decode("utf-8-sig")
            except UnicodeDecodeError: text = content.decode("latin-1")
            rows = [dict(r) for r in csv.DictReader(io.StringIO(text))]
        elif ext in ("xls","xlsx"):
            if not HAS_EXCEL: return jsonify({"error": "openpyxl missing"}), 500
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active; headers = None
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    headers = [str(c).strip().lower() if c else "" for c in row]; continue
                if headers:
                    rows.append({headers[j]: (row[j] if j < len(row) else "") for j in range(len(headers))})
        elif ext == "mdb":
            return jsonify({"error": "MDB import needs MS Access driver."}), 400
        else:
            return jsonify({"error": f"Unsupported: {ext}"}), 400
    except Exception as e:
        return jsonify({"error": f"Parse error: {e}"}), 400

    added = 0
    for r in rows:
        if not isinstance(r, dict): continue
        site = str(r.get("site") or r.get("Site") or r.get("name") or "").strip()
        username = str(r.get("username") or r.get("Username") or "").strip()
        password = str(r.get("password") or r.get("Password") or "")
        url = str(r.get("url") or r.get("URL") or "").strip()
        notes = str(r.get("notes") or r.get("Notes") or "")
        if not any([site, username, password, url, notes]):
            continue
        STATE["data"].append({
            "id": str(uuid.uuid4()), "site": site, "username": username,
            "password": password, "url": url, "notes": notes,
            "category": str(r.get("category") or r.get("Category") or "general").strip() or "general",
            "created": now_iso(), "modified": now_iso(), "password_changed": now_iso(),
        })
        added += 1
    save_vault(); log_audit("IMPORT", f"{ext} ({added} entries)")
    return jsonify({"success": True, "imported": added})

# ================= STATS / HEALTH / AUDIT =================
@app.route("/api/stats", methods=["GET"])
@require_unlock
def stats():
    entries = STATE["data"]
    pw_count = {}
    for e in entries:
        p = e.get("password","")
        if p: pw_count[p] = pw_count.get(p,0)+1
    weak = sum(1 for e in entries if password_strength(e.get("password","")) < 40)
    reused = sum(1 for p,c in pw_count.items() if c > 1)
    cats = {}
    for e in entries:
        c = e.get("category") or "general"
        cats[c] = cats.get(c,0)+1
    threshold = datetime.utcnow() - timedelta(days=90)
    old = 0
    for e in entries:
        try:
            if datetime.fromisoformat(e.get("password_changed","")) < threshold: old += 1
        except Exception: pass
    return jsonify({"total": len(entries), "weak": weak, "reused": reused,
                    "old": old, "categories": cats})

@app.route("/api/health-check", methods=["GET"])
@require_unlock
def health_check():
    entries = STATE["data"]
    pw_count = {}
    for e in entries:
        p = e.get("password","")
        if p: pw_count[p] = pw_count.get(p,0)+1
    weak, reused, old = [], [], []
    threshold = datetime.utcnow() - timedelta(days=90)
    for e in entries:
        pw = e.get("password",""); s = password_strength(pw)
        if s < 40: weak.append({"id": e["id"], "site": e["site"] or e["username"] or "(untitled)", "score": s})
        if pw and pw_count.get(pw,0) > 1:
            reused.append({"id": e["id"], "site": e["site"] or e["username"] or "(untitled)", "count": pw_count[pw]})
        try:
            if datetime.fromisoformat(e.get("password_changed","")) < threshold:
                old.append({"id": e["id"], "site": e["site"] or e["username"] or "(untitled)",
                            "changed": e.get("password_changed","")})
        except Exception: pass
    return jsonify({"weak": weak, "reused": reused, "old": old})

@app.route("/api/audit", methods=["GET"])
@require_unlock
def get_audit():
    if not AUDIT_PATH.exists(): return jsonify({"lines": []})
    lines = AUDIT_PATH.read_text(encoding="utf-8").splitlines()[-300:]
    return jsonify({"lines": list(reversed(lines))})

if __name__ == "__main__":
    print("=" * 62)
    print("  PassManager Professional — Flask Backend")
    print("  → http://127.0.0.1:5000")
    print(f"  → Argon2id: {HAS_ARGON2} | Excel: {HAS_EXCEL} | MDB: {HAS_MDB}")
    print("=" * 62)
    app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)